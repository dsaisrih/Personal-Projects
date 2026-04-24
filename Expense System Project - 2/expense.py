import tkinter as tk
from tkinter import ttk, messagebox
from tkcalendar import Calendar
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import simpleSplit
import mysql.connector
from collections import defaultdict
import matplotlib.pyplot as plt
import os
from datetime import datetime
import tempfile
from io import BytesIO

# ================= COLORS =================
BG = "#f1f5f9"
PRIMARY = "#2563eb"
SUCCESS = "#22c55e"
DANGER = "#ef4444"
CARD = "#e0f2fe"

# ================= DATABASE =================
conn = mysql.connector.connect(
    host="localhost",
    user="root",
    password="D$ai1919"
)
cur = conn.cursor()
cur.execute("CREATE DATABASE IF NOT EXISTS expense_db")
cur.execute("USE expense_db")

cur.execute("""
CREATE TABLE IF NOT EXISTS users(
    id INT AUTO_INCREMENT PRIMARY KEY,
    username VARCHAR(100) UNIQUE,
    password VARCHAR(100)
)
""")

cur.execute("""
CREATE TABLE IF NOT EXISTS expenses(
    id INT AUTO_INCREMENT PRIMARY KEY,
    user_id INT,
    date DATE,
    category VARCHAR(50),
    paid_through VARCHAR(50),
    description VARCHAR(255),
    amount DECIMAL(10,2),
    FOREIGN KEY (user_id) REFERENCES users(id)
)
""")

cur.execute("""
CREATE TABLE IF NOT EXISTS balance(
    user_id INT PRIMARY KEY,
    amount DECIMAL(10,2),
    FOREIGN KEY (user_id) REFERENCES users(id)
)
""")
conn.commit()

# ================= EXPORT FOLDERS =================
BASE = os.getcwd()
PDF_DIR = os.path.join(BASE, "PDF")
EXCEL_DIR = os.path.join(BASE, "EXCEL")
os.makedirs(PDF_DIR, exist_ok=True)
os.makedirs(EXCEL_DIR, exist_ok=True)

# ================= LOGIN =================
login_win = None

def create_login_window():
    global login_win
    login_win = tk.Tk()
    login_win.title("Expense Manager - Login")
    login_win.geometry("500x450")
    login_win.configure(bg=BG)
    login_win.resizable(False, False)
    login_win.grab_set()

    # Header Frame
    header = tk.Frame(login_win, bg=PRIMARY, height=100)
    header.pack(fill=tk.X, padx=0, pady=0)
    header.pack_propagate(False)
    
    tk.Label(header, text="💰 Expense Manager", font=("Segoe UI", 26, "bold"),
             fg="white", bg=PRIMARY).pack(pady=25)

    # Main Content Frame
    content = tk.Frame(login_win, bg=BG)
    content.pack(fill=tk.BOTH, expand=True, padx=40, pady=40)

    # Using Grid for better control
    # Row 0: Username Label
    tk.Label(content, text="Username:", font=("Segoe UI", 11, "bold"),
             fg=PRIMARY, bg=BG).grid(row=0, column=0, sticky="w", pady=(0, 5))
    
    # Row 0, Col 1: Username Entry
    u_entry = tk.Entry(content, font=("Segoe UI", 11), width=25,
                       relief=tk.SOLID, bd=1, bg="white")
    u_entry.grid(row=0, column=1, sticky="ew", pady=(0, 15))
    
    # Row 1: Password Label
    tk.Label(content, text="Password:", font=("Segoe UI", 11, "bold"),
             fg=PRIMARY, bg=BG).grid(row=1, column=0, sticky="w", pady=(0, 5))
    
    # Row 1, Col 1: Password Entry
    p_entry = tk.Entry(content, font=("Segoe UI", 11), width=25,
                       relief=tk.SOLID, bd=1, bg="white", show="•")
    p_entry.grid(row=1, column=1, sticky="ew", pady=(0, 30))
    
    # Configure grid columns
    content.columnconfigure(1, weight=1)

    def login():
        username_val = u_entry.get().strip()
        password_val = p_entry.get()
        
        if not username_val:
            messagebox.showerror("Error", "Please enter username")
            u_entry.focus()
            return
        
        if not password_val:
            messagebox.showerror("Error", "Please enter password")
            p_entry.focus()
            return
            
        try:
            cur.execute(
                "SELECT id, username FROM users WHERE username=%s AND password=%s",
                (username_val, password_val)
            )
            result = cur.fetchone()
            if result:
                user_id, username = result
                login_win.withdraw()
                main_app(user_id, username, login_win)
            else:
                messagebox.showerror("Error", "Invalid username or password")
                p_entry.delete(0, tk.END)
                u_entry.focus()
        except Exception as e:
            messagebox.showerror("Error", f"Login failed: {str(e)}")

    def register():
        username_val = u_entry.get().strip()
        password_val = p_entry.get()
        
        if not username_val:
            messagebox.showerror("Error", "Please enter username")
            u_entry.focus()
            return
        
        if not password_val:
            messagebox.showerror("Error", "Please enter password")
            p_entry.focus()
            return
        
        if len(username_val) < 3:
            messagebox.showerror("Error", "Username must be at least 3 characters")
            u_entry.focus()
            return
        
        if len(password_val) < 4:
            messagebox.showerror("Error", "Password must be at least 4 characters")
            p_entry.focus()
            return
            
        try:
            cur.execute(
                "INSERT INTO users VALUES(NULL,%s,%s)",
                (username_val, password_val)
            )
            conn.commit()
            messagebox.showinfo("Success", "Account Created!\nNow login with your credentials")
            u_entry.delete(0, tk.END)
            p_entry.delete(0, tk.END)
            u_entry.focus()
        except Exception as e:
            if "Duplicate" in str(e):
                messagebox.showerror("Error", "Username already exists!")
                u_entry.delete(0, tk.END)
                u_entry.focus()
            else:
                messagebox.showerror("Error", str(e))

    # Row 2: Button Frame
    btn_frame = tk.Frame(content, bg=BG)
    btn_frame.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(10, 0))

    login_btn = tk.Button(btn_frame, text="🔓 Login", bg=SUCCESS, fg="white",
                          font=("Segoe UI", 10, "bold"), command=login,
                          relief=tk.RAISED, bd=1, width=15)
    login_btn.pack(side=tk.LEFT, padx=5)

    register_btn = tk.Button(btn_frame, text="📝 Register", bg=PRIMARY, fg="white",
                             font=("Segoe UI", 10, "bold"), command=register,
                             relief=tk.RAISED, bd=1, width=15)
    register_btn.pack(side=tk.LEFT, padx=5)

    # Bind Enter key to login
    u_entry.bind("<Return>", lambda e: login())
    p_entry.bind("<Return>", lambda e: login())

    u_entry.focus()
    login_win.mainloop()

# ================= MAIN APP =================
def main_app(user_id, username, login_window):
    root = tk.Tk()
    root.title("Expense Management System")
    root.geometry("1350x850")
    root.configure(bg=BG)

    selected_id = tk.IntVar()

    # ---------- BALANCE ----------
    def load_balance():
        cur.execute("SELECT amount FROM balance WHERE user_id=%s", (user_id,))
        result = cur.fetchone()
        balance = result[0] if result else 0
        bal_label.config(
            text=f"Balance : ₹ {balance:,.2f}"
        )

    def update_balance():
        cur.execute("SELECT COUNT(*) FROM balance WHERE user_id=%s", (user_id,))
        if cur.fetchone()[0] == 0:
            cur.execute("INSERT INTO balance VALUES(%s, %s)", (user_id, bal_entry.get()))
        else:
            cur.execute(
                "UPDATE balance SET amount = amount + %s WHERE user_id=%s",
                (bal_entry.get(), user_id)
            )
        conn.commit()
        bal_entry.delete(0, tk.END)
        load_balance()

    # ---------- DATE PICKER ----------
    def pick_date():
        def select():
            m, d, y = cal.get_date().split("/")
            date_entry.delete(0, tk.END)
            date_entry.insert(
                0, f"{y}-{int(m):02d}-{int(d):02d}"
            )
            win.destroy()

        win = tk.Toplevel(root)
        cal = Calendar(win)
        cal.pack()
        tk.Button(win, text="Select",
                  command=select).pack()

    # ---------- CRUD ----------
    def add_expense():
        # Get all values
        date_val = date_entry.get().strip()
        cat_val = cat_combo.get().strip()
        paid_val = paid_combo.get().strip()
        desc_val = desc.get().strip()
        amt_val = amt.get().strip()
        
        # Validate each field
        if not date_val:
            messagebox.showerror("Error", "Please select a date")
            date_entry.focus()
            return
        if not cat_val:
            messagebox.showerror("Error", "Please select a category")
            return
        if not paid_val:
            messagebox.showerror("Error", "Please select payment method")
            return
        if not desc_val:
            messagebox.showerror("Error", "Please enter description")
            desc.focus()
            return
        if not amt_val:
            messagebox.showerror("Error", "Please enter amount")
            amt.focus()
            return
            
        try:
            amount = float(amt_val)
            if amount <= 0:
                messagebox.showerror("Error", "Amount must be greater than 0")
                amt.focus()
                return
                
            cur.execute(
                "INSERT INTO expenses VALUES(NULL,%s,%s,%s,%s,%s,%s)",
                (user_id, date_val, cat_val, paid_val, desc_val, amount)
            )
            cur.execute(
                "UPDATE balance SET amount = amount - %s WHERE user_id=%s",
                (amount, user_id)
            )
            conn.commit()
            load_expenses()
            load_balance()
            date_entry.delete(0, tk.END)
            cat_combo.set("Food")
            paid_combo.set("Cash")
            desc.delete(0, tk.END)
            amt.delete(0, tk.END)
            messagebox.showinfo("Success", "Expense added successfully")
        except ValueError:
            messagebox.showerror("Error", "Amount must be a valid number")
            amt.focus()
        except Exception as e:
            messagebox.showerror("Error", f"Error adding expense: {str(e)}")

    def load_expenses():
        tree.delete(*tree.get_children())
        cur.execute("SELECT * FROM expenses WHERE user_id=%s ORDER BY date", (user_id,))
        for r in cur.fetchall():
            # Display: ID, Date, Category, Paid Through, Description, Amount (skip user_id)
            tree.insert("", tk.END, values=(r[0], r[2], r[3], r[4], r[5], r[6]))

    def select_row(e):
        row = tree.item(tree.focus())["values"]
        if not row:
            return
        selected_id.set(row[0])  # id
        date_entry.delete(0, tk.END)
        date_entry.insert(0, row[1])  # date
        
        # Set category - make sure it matches exactly
        cat_value = row[2]  # category
        cat_combo.set(cat_value)
        
        # Set payment method - make sure it matches exactly
        paid_value = row[3]  # paid_through
        paid_combo.set(paid_value)
        
        desc.delete(0, tk.END)
        desc.insert(0, row[4])  # description
        amt.delete(0, tk.END)
        amt.insert(0, row[5])  # amount

    def update_expense():
        if not selected_id.get():
            messagebox.showwarning("Warning", "Please select an expense to update")
            return
        try:
            # Get old amount to calculate difference
            cur.execute("SELECT amount FROM expenses WHERE id=%s AND user_id=%s", (selected_id.get(), user_id))
            old_result = cur.fetchone()
            if not old_result:
                messagebox.showerror("Error", "Expense not found")
                return
            old_amount = float(old_result[0])
            new_amount = float(amt.get())
            amount_diff = new_amount - old_amount
            
            # Update expense
            cur.execute(
                """UPDATE expenses SET date=%s, category=%s,
                   paid_through=%s, description=%s,
                   amount=%s WHERE id=%s AND user_id=%s""",
                (date_entry.get(), cat_combo.get(),
                 paid_combo.get(), desc.get(),
                 new_amount, selected_id.get(), user_id)
            )
            # Adjust balance by the difference
            cur.execute(
                "UPDATE balance SET amount = amount - %s WHERE user_id=%s",
                (amount_diff, user_id)
            )
            conn.commit()
            load_expenses()
            load_balance()
            messagebox.showinfo("Success", "Expense updated successfully")
        except Exception as e:
            messagebox.showerror("Error", f"Error updating expense: {str(e)}")

    def delete_expense():
        if not selected_id.get():
            messagebox.showwarning("Warning", "Please select an expense to delete")
            return
        if messagebox.askyesno("Confirm", "Are you sure you want to delete this expense?"):
            try:
                # Get amount to refund
                cur.execute("SELECT amount FROM expenses WHERE id=%s AND user_id=%s", (selected_id.get(), user_id))
                result = cur.fetchone()
                if result:
                    amount = float(result[0])
                    # Delete expense
                    cur.execute(
                        "DELETE FROM expenses WHERE id=%s AND user_id=%s",
                        (selected_id.get(), user_id)
                    )
                    # Refund balance
                    cur.execute(
                        "UPDATE balance SET amount = amount + %s WHERE user_id=%s",
                        (amount, user_id)
                    )
                    conn.commit()
                    load_expenses()
                    load_balance()
                    messagebox.showinfo("Success", "Expense deleted and balance refunded")
            except Exception as e:
                messagebox.showerror("Error", f"Error deleting expense: {str(e)}")

    # ---------- CHARTS ----------
    def bar_chart():
        data = defaultdict(float)
        cur.execute("SELECT date, amount FROM expenses WHERE user_id=%s", (user_id,))
        for d, a in cur.fetchall():
            data[d.strftime("%Y-%m")] += float(a)
        plt.bar(data.keys(), data.values())
        plt.xticks(rotation=45)
        plt.title("Monthly Expenses")
        plt.tight_layout()
        plt.show()

    def pie_chart():
        data = defaultdict(float)
        cur.execute("SELECT category, amount FROM expenses WHERE user_id=%s", (user_id,))
        for c, a in cur.fetchall():
            data[c] += float(a)
        plt.pie(data.values(),
                labels=data.keys(),
                autopct="%1.1f%%")
        plt.title("Expenses by Category")
        plt.show()

    # ---------- EXPORT DATE ----------
    def today_str():
        return datetime.now().strftime("%d-%m-%Y")
    
    def generate_chart_images():
        # -------- BAR CHART (Monthly) --------
        data_month = defaultdict(float)
        cur.execute("SELECT date, amount FROM expenses WHERE user_id=%s", (user_id,))
        for d, a in cur.fetchall():
            data_month[d.strftime("%Y-%m")] += float(a)

        # Save to temporary file
        bar_path = os.path.join(tempfile.gettempdir(), "bar_chart_temp.png")
        plt.figure(figsize=(6, 3))
        plt.bar(data_month.keys(), data_month.values())
        plt.xticks(rotation=45)
        plt.title("Monthly Expenses")
        plt.tight_layout()
        plt.savefig(bar_path)
        plt.close()

        # -------- PIE CHART (Category) --------
        data_cat = defaultdict(float)
        cur.execute("SELECT category, amount FROM expenses WHERE user_id=%s", (user_id,))
        for c, a in cur.fetchall():
            data_cat[c] += float(a)

        # Save to temporary file
        pie_path = os.path.join(tempfile.gettempdir(), "pie_chart_temp.png")
        plt.figure(figsize=(5, 5))
        plt.pie(data_cat.values(), labels=data_cat.keys(), autopct="%1.1f%%")
        plt.title("Expenses by Category")
        plt.tight_layout()
        plt.savefig(pie_path)
        plt.close()

        return bar_path, pie_path


    # ---------- EXPORT PDF ----------
    def export_pdf():
        cur.execute(
            "SELECT date, category, paid_through, description, amount FROM expenses WHERE user_id=%s ORDER BY date",
            (user_id,)
        )
        rows = cur.fetchall()

        if not rows:
            messagebox.showinfo("No Data", "No expenses to export")
            return

        filename = f"Expense_Report_{today_str()}.pdf"
        file_path = os.path.join(PDF_DIR, filename)

        # Generate charts
        bar_img, pie_img = generate_chart_images()

        c = canvas.Canvas(file_path, pagesize=A4)
        w, h = A4

        # -------- TITLE --------
        c.setFont("Helvetica-Bold", 18)
        c.drawString(40, h - 40, "EXPENSE REPORT")

        # -------- TABLE HEADER --------
        headers = ["DATE", "CATEGORY", "PAID THROUGH", "DESCRIPTION", "AMOUNT"]
        x = [40, 100, 200, 300, 480]
        y = h - 80

        c.setFont("Helvetica-Bold", 9)
        for i in range(len(headers)):
            c.drawString(x[i], y, headers[i])

        c.setFont("Helvetica", 9)
        y -= 20

        total_amount = 0
        for r in rows:
            total_amount += float(r[4])
            desc_lines = simpleSplit(r[3], "Helvetica", 9, 180)

            c.drawString(x[0], y, r[0].strftime("%d-%m-%Y"))
            c.drawString(x[1], y, r[1])
            c.drawString(x[2], y, r[2])
            c.drawString(x[4], y, f"₹ {r[4]:,.2f}")

            for line in desc_lines:
                c.drawString(x[3], y, line)
                y -= 14
                if y < 60:
                    c.showPage()
                    y = h - 60
                    c.setFont("Helvetica", 9)

        # -------- TOTAL --------
        y -= 10
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x[3], y, "TOTAL")
        c.drawString(x[4], y, f"₹ {total_amount:,.2f}")

        # -------- NEW PAGE FOR CHARTS --------
        c.showPage()
        c.setFont("Helvetica-Bold", 16)
        c.drawString(40, h - 40, "Expense Analysis")

        # Bar chart (top)
        c.drawString(40, h - 80, "Monthly Expenses:")
        c.drawImage(bar_img, 40, h - 350, width=500, height=250)

        # Pie chart (bottom)
        c.drawString(40, h - 380, "Expenses by Category:")
        c.drawImage(pie_img, 120, h - 700, width=350, height=300)

        c.save()
        
        # Clean up temporary files
        try:
            os.remove(bar_img)
            os.remove(pie_img)
        except:
            pass
        
        messagebox.showinfo("PDF", f"PDF Exported:\n{filename}")

    

    # ---------- EXPORT EXCEL ----------
    def export_excel():
        cur.execute(
            "SELECT * FROM expenses WHERE user_id=%s ORDER BY date",
            (user_id,)
        )
        rows = cur.fetchall()

        wb = Workbook()
        ws = wb.active
        ws.append(["ID","Date","Category","Paid Through","Description","Amount"])

        total_amount = 0
        for r in rows:
            total_amount += float(r[5])
            ws.append(r)
        ws.append(["","","","TOTAL","","{:.2f}".format(total_amount)])

        filename = f"Expense_Report_{today_str()}.xlsx"
        file_path = os.path.join(EXCEL_DIR, filename)
        wb.save(file_path)
        messagebox.showinfo("Excel", f"Excel Exported:\n{filename}")

    # ---------- GENERATE FILTERED CHARTS ----------
    def generate_filtered_chart_images(month, year):
        # -------- BAR CHART (Monthly) --------
        data_month = defaultdict(float)
        if month == 0:
            cur.execute("SELECT date, amount FROM expenses WHERE YEAR(date)=%s AND user_id=%s", (year, user_id))
        else:
            cur.execute("SELECT date, amount FROM expenses WHERE MONTH(date)=%s AND YEAR(date)=%s AND user_id=%s", (month, year, user_id))
        
        for d, a in cur.fetchall():
            data_month[d.strftime("%Y-%m")] += float(a)

        if not data_month:
            return None, None

        # Save to temporary file
        bar_path = os.path.join(tempfile.gettempdir(), "bar_chart_temp.png")
        plt.figure(figsize=(6, 3))
        plt.bar(data_month.keys(), data_month.values())
        plt.xticks(rotation=45)
        plt.title("Monthly Expenses")
        plt.tight_layout()
        plt.savefig(bar_path)
        plt.close()

        # -------- PIE CHART (Category) --------
        data_cat = defaultdict(float)
        if month == 0:
            cur.execute("SELECT category, amount FROM expenses WHERE YEAR(date)=%s AND user_id=%s", (year, user_id))
        else:
            cur.execute("SELECT category, amount FROM expenses WHERE MONTH(date)=%s AND YEAR(date)=%s AND user_id=%s", (month, year, user_id))
        
        for c, a in cur.fetchall():
            data_cat[c] += float(a)

        # Save to temporary file
        pie_path = os.path.join(tempfile.gettempdir(), "pie_chart_temp.png")
        plt.figure(figsize=(5, 5))
        plt.pie(data_cat.values(), labels=data_cat.keys(), autopct="%1.1f%%")
        plt.title("Expenses by Category")
        plt.tight_layout()
        plt.savefig(pie_path)
        plt.close()

        return bar_path, pie_path

    # ---------- EXPORT BY MONTH & YEAR ----------
    def export_by_month_year():
        month = month_entry.get()
        year = year_entry.get()

        if not month.isdigit() or not year.isdigit():
            messagebox.showerror("Error", "Enter valid Month and Year")
            return

        month = int(month)
        year = int(year)

        if month == 0:
            cur.execute(
                "SELECT date, category, paid_through, description, amount "
                "FROM expenses WHERE YEAR(date)=%s AND user_id=%s ORDER BY date",
                (year, user_id)
            )
            filename = f"Expense_Report_Year_{year}.pdf"
        else:
            cur.execute(
                "SELECT date, category, paid_through, description, amount "
                "FROM expenses WHERE MONTH(date)=%s AND YEAR(date)=%s AND user_id=%s ORDER BY date",
                (month, year, user_id)
            )
            filename = f"Expense_Report_{month}_{year}.pdf"

        rows = cur.fetchall()
        if not rows:
            messagebox.showinfo("No Data", "No expenses found for this period")
            return

        # Generate filtered charts
        bar_img, pie_img = generate_filtered_chart_images(month, year)

        # ---------------- PDF ----------------
        file_path = os.path.join(PDF_DIR, filename)
        c = canvas.Canvas(file_path, pagesize=A4)
        w, h = A4

        c.setFont("Helvetica-Bold", 18)
        title = f"EXPENSE REPORT - {year}" if month==0 else f"EXPENSE REPORT - {month}/{year}"
        c.drawString(40, h - 40, title)

        headers = ["DATE", "CATEGORY", "PAID THROUGH", "DESCRIPTION", "AMOUNT"]
        x = [40, 100, 200, 300, 480]
        y = h - 80

        c.setFont("Helvetica-Bold", 9)
        for i in range(len(headers)):
            c.drawString(x[i], y, headers[i])

        c.setFont("Helvetica", 9)
        y -= 20

        total_amount = 0
        for r in rows:
            total_amount += float(r[4])
            desc_lines = simpleSplit(r[3], "Helvetica", 9, 180)
            c.drawString(x[0], y, r[0].strftime("%d-%m-%Y"))
            c.drawString(x[1], y, r[1])
            c.drawString(x[2], y, r[2])
            c.drawString(x[4], y, f"₹ {r[4]:,.2f}")
            for line in desc_lines:
                c.drawString(x[3], y, line)
                y -= 14
                if y < 50:
                    c.showPage()
                    c.setFont("Helvetica", 9)
                    y = h - 60

        y -= 10
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x[3], y, "TOTAL")
        c.drawString(x[4], y, f"₹ {total_amount:,.2f}")

        # Add charts page
        if bar_img and pie_img:
            c.showPage()
            c.setFont("Helvetica-Bold", 16)
            c.drawString(40, h - 40, "Expense Analysis")

            # Bar chart (top)
            c.drawString(40, h - 80, "Monthly Expenses:")
            c.drawImage(bar_img, 40, h - 350, width=500, height=250)

            # Pie chart (bottom)
            c.drawString(40, h - 380, "Expenses by Category:")
            c.drawImage(pie_img, 120, h - 700, width=350, height=300)

        c.save()

        # ---------------- EXCEL ----------------
        filename_xl = filename.replace(".pdf", ".xlsx")
        file_path_xl = os.path.join(EXCEL_DIR, filename_xl)
        wb = Workbook()
        ws = wb.active
        ws.append(headers)

        for r in rows:
            ws.append(r)
        ws.append(["","","","TOTAL","","{:.2f}".format(total_amount)])

        # Add charts to Excel
        if bar_img and pie_img:
            # Add blank rows before charts
            ws.append([])
            ws.append(["Expense Analysis"])
            
            # Add bar chart
            ws.append(["Monthly Expenses:"])
            bar_img_xl = XLImage(bar_img)
            bar_img_xl.width = 500
            bar_img_xl.height = 250
            ws.add_image(bar_img_xl, 'A' + str(ws.max_row + 1))
            
            # Move to next chart position
            for _ in range(5):
                ws.append([])
            
            # Add pie chart
            ws.append(["Expenses by Category:"])
            pie_img_xl = XLImage(pie_img)
            pie_img_xl.width = 350
            pie_img_xl.height = 300
            ws.add_image(pie_img_xl, 'A' + str(ws.max_row + 1))

        wb.save(file_path_xl)

        # Clean up temporary files - AFTER both PDF and Excel are created
        try:
            if bar_img and pie_img:
                os.remove(bar_img)
                os.remove(pie_img)
        except:
            pass

        messagebox.showinfo("PDF & Excel", f"PDF and Excel Exported:\n{filename}")

    # ---------- LOGOUT ----------
    def logout():
        if messagebox.askyesno("Logout", "Are you sure you want to logout?"):
            root.destroy()
            login_window.deiconify()
            create_login_window()

    # ---------- UI ----------
    # Header Frame
    header_frame = tk.Frame(root, bg=PRIMARY, height=80)
    header_frame.pack(fill=tk.X, padx=0, pady=0)
    header_frame.pack_propagate(False)
    
    # Title and User info in header
    title_user_frame = tk.Frame(header_frame, bg=PRIMARY)
    title_user_frame.pack(fill=tk.X, pady=15, padx=20)
    
    tk.Label(title_user_frame, text="💰 Expense Management System", 
             font=("Segoe UI", 28, "bold"), fg="white", bg=PRIMARY).pack(side=tk.LEFT)
    
    tk.Label(title_user_frame, text=f"👤 {username}", 
             font=("Segoe UI", 12, "bold"), fg="#e0f2fe", bg=PRIMARY).pack(side=tk.RIGHT, padx=(0, 20))
    
    tk.Button(title_user_frame, text="🚪 Logout", bg=DANGER, fg="white",
              font=("Segoe UI", 9, "bold"), command=logout).pack(side=tk.RIGHT, padx=5)

    # Balance Frame
    bal_frame = tk.Frame(root, bg=CARD, height=70)
    bal_frame.pack(fill=tk.X, padx=15, pady=15)
    bal_frame.pack_propagate(False)
    
    tk.Label(bal_frame, text="Current Balance", font=("Segoe UI", 11, "bold"), 
             bg=CARD, fg=PRIMARY).pack(side=tk.LEFT, padx=15, pady=10)
    
    bal_label = tk.Label(bal_frame, text="Balance : ₹ 0.00", 
                         bg=CARD, font=("Segoe UI", 14, "bold"), fg=SUCCESS)
    bal_label.pack(side=tk.LEFT, padx=20)
    
    tk.Label(bal_frame, text="Add Amount:", font=("Segoe UI", 10), 
             bg=CARD).pack(side=tk.LEFT, padx=10)
    
    bal_entry = tk.Entry(bal_frame, width=15, font=("Segoe UI", 10))
    bal_entry.pack(side=tk.LEFT, padx=5)
    
    tk.Button(bal_frame, text="✓ Add Balance", bg=SUCCESS, fg="white", 
              font=("Segoe UI", 9, "bold"), command=update_balance).pack(side=tk.LEFT, padx=5)

    # Main content container
    main_container = tk.Frame(root, bg=BG)
    main_container.pack(fill=tk.BOTH, expand=True, padx=15, pady=10)

    # Left side - Form
    left_frame = tk.Frame(main_container, bg=BG)
    left_frame.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 10))

    # Add Expense Section
    form_header = tk.Label(left_frame, text="Add/Update Expense", 
                           font=("Segoe UI", 12, "bold"), fg=PRIMARY, bg=BG)
    form_header.pack(pady=(0, 10))

    form = tk.Frame(left_frame, bg="white", relief=tk.RAISED, bd=1)
    form.pack(padx=10, pady=10, fill=tk.X)

    # Helper function for styled labels
    def create_form_row(parent, label_text, widget):
        row_frame = tk.Frame(parent, bg="white")
        row_frame.pack(fill=tk.X, padx=10, pady=8)
        tk.Label(row_frame, text=label_text, font=("Segoe UI", 9, "bold"), 
                 bg="white", fg="#374151", width=15).pack(side=tk.LEFT)
        widget.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

    # Date row
    date_row = tk.Frame(form, bg="white")
    date_row.pack(fill=tk.X, padx=10, pady=8)
    tk.Label(date_row, text="Date", font=("Segoe UI", 9, "bold"), 
             bg="white", fg="#374151", width=15).pack(side=tk.LEFT)
    date_entry = tk.Entry(date_row, width=20, font=("Segoe UI", 9))
    date_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))
    tk.Button(date_row, text="📅 Pick", bg=PRIMARY, fg="white", 
              font=("Segoe UI", 8), command=pick_date, width=8).pack(side=tk.LEFT)

    # Category row
    cat_row = tk.Frame(form, bg="white")
    cat_row.pack(fill=tk.X, padx=10, pady=8)
    tk.Label(cat_row, text="Category", font=("Segoe UI", 9, "bold"), 
             bg="white", fg="#374151", width=15).pack(side=tk.LEFT)
    cat_combo = ttk.Combobox(cat_row, 
                             values=["Food", "Transport", "Rent", "Shopping", "Bills", "Entertainment", "Other"],
                             font=("Segoe UI", 9), state="readonly", width=23)
    cat_combo.set("   ")
    cat_combo.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

    # Payment Method row
    paid_row = tk.Frame(form, bg="white")
    paid_row.pack(fill=tk.X, padx=10, pady=8)
    tk.Label(paid_row, text="Paid Through", font=("Segoe UI", 9, "bold"), 
             bg="white", fg="#374151", width=15).pack(side=tk.LEFT)
    paid_combo = ttk.Combobox(paid_row, 
                              values=["Cash", "Credit", "UPI", "Other"],
                              font=("Segoe UI", 9), state="readonly", width=23)
    paid_combo.set("   ")
    paid_combo.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

    # Description row
    desc_row = tk.Frame(form, bg="white")
    desc_row.pack(fill=tk.X, padx=10, pady=8)
    tk.Label(desc_row, text="Description", font=("Segoe UI", 9, "bold"), 
             bg="white", fg="#374151", width=15).pack(side=tk.LEFT)
    desc = tk.Entry(desc_row, width=27, font=("Segoe UI", 9))
    desc.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

    # Amount row
    amt_row = tk.Frame(form, bg="white")
    amt_row.pack(fill=tk.X, padx=10, pady=8)
    tk.Label(amt_row, text="Amount (₹)", font=("Segoe UI", 9, "bold"), 
             bg="white", fg="#374151", width=15).pack(side=tk.LEFT)
    amt = tk.Entry(amt_row, width=27, font=("Segoe UI", 9))
    amt.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

    # Action Buttons
    action_frame = tk.Frame(form, bg="white")
    action_frame.pack(fill=tk.X, padx=10, pady=12)
    
    tk.Button(action_frame, text="✓ Add", bg=SUCCESS, fg="white", 
              font=("Segoe UI", 9, "bold"), command=add_expense, width=8).pack(side=tk.LEFT, padx=2)
    tk.Button(action_frame, text="✎ Update", bg=PRIMARY, fg="white", 
              font=("Segoe UI", 9, "bold"), command=update_expense, width=8).pack(side=tk.LEFT, padx=2)
    tk.Button(action_frame, text="✕ Delete", bg=DANGER, fg="white", 
              font=("Segoe UI", 9, "bold"), command=delete_expense, width=8).pack(side=tk.LEFT, padx=2)

    # Filter Section
    filter_header = tk.Label(left_frame, text="Report by Month-Year", 
                            font=("Segoe UI", 12, "bold"), fg=PRIMARY, bg=BG)
    filter_header.pack(pady=(15, 10))

    filter_frame = tk.Frame(left_frame, bg="white", relief=tk.RAISED, bd=1)
    filter_frame.pack(padx=10, pady=10, fill=tk.X)

    # Month row
    month_row = tk.Frame(filter_frame, bg="white")
    month_row.pack(fill=tk.X, padx=10, pady=8)
    tk.Label(month_row, text="Month (MM)", font=("Segoe UI", 9, "bold"), 
             bg="white", fg="#374151", width=15).pack(side=tk.LEFT)
    month_entry = tk.Entry(month_row, width=27, font=("Segoe UI", 9))
    month_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))
    tk.Label(month_row, text="00 = Full Year", font=("Segoe UI", 8), 
             fg="#6b7280", bg="white").pack(side=tk.LEFT)

    # Year row
    year_row = tk.Frame(filter_frame, bg="white")
    year_row.pack(fill=tk.X, padx=10, pady=8)
    tk.Label(year_row, text="Year (YYYY)", font=("Segoe UI", 9, "bold"), 
             bg="white", fg="#374151", width=15).pack(side=tk.LEFT)
    year_entry = tk.Entry(year_row, width=27, font=("Segoe UI", 9))
    year_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

    # Export buttons
    export_frame = tk.Frame(filter_frame, bg="white")
    export_frame.pack(fill=tk.X, padx=10, pady=10)
    tk.Button(export_frame, text="📅 Generate Report", bg=PRIMARY, fg="white", 
              font=("Segoe UI", 9, "bold"), command=export_by_month_year, width=30).pack(fill=tk.X)

    # Right side - Data Table and Charts
    right_frame = tk.Frame(main_container, bg=BG)
    right_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

    # Expenses Table Header
    table_header = tk.Label(right_frame, text="Expenses Records", 
                           font=("Segoe UI", 12, "bold"), fg=PRIMARY, bg=BG)
    table_header.pack(pady=(0, 10))

    # Treeview with custom styling
    tree = ttk.Treeview(right_frame, columns=("ID","Date","Category","Paid Through","Description","Amount"), 
                       show="headings", height=10)
    tree.column("ID", width=30)
    tree.column("Date", width=80)
    tree.column("Category", width=80)
    tree.column("Paid Through", width=90)
    tree.column("Description", width=150)
    tree.column("Amount", width=80)

    for c in tree["columns"]:
        tree.heading(c, text=c)
    
    tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
    tree.bind("<ButtonRelease-1>", select_row)

    # Charts and Export Section
    btn_frame = tk.Frame(right_frame, bg=BG)
    btn_frame.pack(fill=tk.X, padx=10, pady=10)

    tk.Label(btn_frame, text="Export & Analytics", font=("Segoe UI", 10, "bold"), 
             fg=PRIMARY, bg=BG).pack(anchor=tk.W, pady=(0, 8))

    # Top row of buttons
    btn_row1 = tk.Frame(btn_frame, bg=BG)
    btn_row1.pack(fill=tk.X, pady=4)
    
    tk.Button(btn_row1, text="📊 Bar Chart", bg="#f59e0b", fg="white", 
              font=("Segoe UI", 9, "bold"), command=bar_chart).pack(side=tk.LEFT, padx=3, fill=tk.X, expand=True)
    tk.Button(btn_row1, text="🥧 Pie Chart", bg="#8b5cf6", fg="white", 
              font=("Segoe UI", 9, "bold"), command=pie_chart).pack(side=tk.LEFT, padx=3, fill=tk.X, expand=True)

    # Bottom row of buttons
    btn_row2 = tk.Frame(btn_frame, bg=BG)
    btn_row2.pack(fill=tk.X, pady=4)
    
    tk.Button(btn_row2, text="📄 PDF (All)", bg="#06b6d4", fg="white", 
              font=("Segoe UI", 9, "bold"), command=export_pdf).pack(side=tk.LEFT, padx=3, fill=tk.X, expand=True)
    tk.Button(btn_row2, text="📊 Excel (All)", bg="#10b981", fg="white", 
              font=("Segoe UI", 9, "bold"), command=export_excel).pack(side=tk.LEFT, padx=3, fill=tk.X, expand=True)

    load_balance()
    load_expenses()
    root.mainloop()

create_login_window()
